from flask import Flask, render_template, request, redirect, url_for, send_file, session, flash
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.page import PageMargins
from openpyxl.utils import get_column_letter
from io import BytesIO
from datetime import date
from werkzeug.security import generate_password_hash, check_password_hash
import sqlite3
import os
import re

app = Flask(__name__)

# =========================================================
# SECURITY / SESSION
# =========================================================

app.secret_key = os.environ.get(
    "SECRET_KEY",
    "development-secret-key-change-in-render"
)

app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = (
    os.environ.get("COOKIE_SECURE", "false").lower() == "true"
)

DATABASE = os.path.join("database", "fmea.db")


# =========================================================
# DATABASE
# =========================================================

def get_db():
    os.makedirs("database", exist_ok=True)
    conn = sqlite3.connect(DATABASE)
    conn.row_factory = sqlite3.Row
    return conn


def add_column_if_missing(cursor, table_name, column_name, column_definition):
    columns = cursor.execute(
        f"PRAGMA table_info({table_name})"
    ).fetchall()

    existing = [column[1] for column in columns]

    if column_name not in existing:
        cursor.execute(
            f"ALTER TABLE {table_name} ADD COLUMN {column_name} {column_definition}"
        )


def setup_database():
    conn = get_db()
    cursor = conn.cursor()

    # USERS
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL UNIQUE,
            password_hash TEXT NOT NULL,
            created_date TEXT NOT NULL
        )
    """)

    admin_username = os.environ.get("ADMIN_USERNAME")
    admin_password = os.environ.get("ADMIN_PASSWORD")

    if admin_username and admin_password:
        existing = cursor.execute(
            "SELECT id FROM users WHERE username = ?",
            (admin_username,)
        ).fetchone()

        if existing is None:
            cursor.execute("""
                INSERT INTO users
                (username, password_hash, created_date)
                VALUES (?, ?, ?)
            """, (
                admin_username,
                generate_password_hash(admin_password),
                date.today().isoformat()
            ))

    # Add registration profile fields to existing databases
    for column, definition in [
        ("name", "TEXT DEFAULT ''"),
        ("organisation", "TEXT DEFAULT ''"),
        ("language", "TEXT DEFAULT 'English'")
    ]:
        add_column_if_missing(cursor, "users", column, definition)

    # PROJECTS
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS projects (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_name TEXT,
            product_name TEXT,
            customer TEXT,
            project_number TEXT,
            created_date TEXT
        )
    """)

    # Each project belongs to exactly one registered user.
    add_column_if_missing(cursor, "projects", "user_id", "INTEGER")

    # Existing projects created before user isolation are assigned to the
    # oldest registered user so existing local data is not lost.
    cursor.execute("""
        UPDATE projects
        SET user_id = (SELECT MIN(id) FROM users)
        WHERE user_id IS NULL
          AND EXISTS (SELECT 1 FROM users)
    """)

    for column, definition in [
        ("project_name", "TEXT"),
        ("product_name", "TEXT"),
        ("customer", "TEXT"),
        ("project_number", "TEXT"),
        ("created_date", "TEXT"),
        ("oem_name", "TEXT DEFAULT 'Generic'"),
        ("compliance_mode", "TEXT DEFAULT 'AIAG-VDA 2019'")
    ]:
        add_column_if_missing(cursor, "projects", column, definition)

    # OEM STANDARDS
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS oem_standards (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            oem_name TEXT NOT NULL UNIQUE,
            standard_framework TEXT NOT NULL,
            cc_symbol TEXT,
            sc_symbol TEXT,
            archiving_period_years INTEGER DEFAULT 10,
            description TEXT
        )
    """)

    oem_data = [
        ("Volkswagen Group", "AIAG-VDA", "D/TLD", "K", 15,
         "OEM CSR prototype configuration for Volkswagen Group."),
        ("BMW Group", "AIAG-VDA", "DS", "PTC", 12,
         "OEM CSR prototype configuration for BMW Group."),
        ("Ford Motor Co", "AIAG-VDA", "∇", "SC", 10,
         "OEM CSR prototype configuration for Ford Motor Co."),
        ("General Motors", "AIAG-VDA", "KPC", "PQC", 10,
         "OEM CSR prototype configuration for General Motors."),
        ("Stellantis", "AIAG-VDA", "S", "R", 10,
         "OEM CSR prototype configuration for Stellantis."),
        ("Generic", "AIAG-VDA 2019", "CC", "SC", 10,
         "Generic FMEA configuration.")
    ]

    cursor.executemany("""
        INSERT OR IGNORE INTO oem_standards
        (oem_name, standard_framework, cc_symbol, sc_symbol,
         archiving_period_years, description)
        VALUES (?, ?, ?, ?, ?, ?)
    """, oem_data)

    # FUNCTIONAL ANALYSIS
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS functional_analysis (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER,
            function TEXT,
            requirement TEXT
        )
    """)

    for column, definition in [
        ("project_id", "INTEGER"),
        ("function", "TEXT"),
        ("requirement", "TEXT")
    ]:
        add_column_if_missing(cursor, "functional_analysis", column, definition)

    # BOUNDARY DIAGRAM
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS boundary_diagram (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER,
            external_element TEXT,
            interaction TEXT,
            direction TEXT,
            description TEXT
        )
    """)

    for column, definition in [
        ("project_id", "INTEGER"),
        ("external_element", "TEXT"),
        ("interaction", "TEXT"),
        ("direction", "TEXT"),
        ("description", "TEXT")
    ]:
        add_column_if_missing(cursor, "boundary_diagram", column, definition)

    # PRODUCT STRUCTURE
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS product_structure (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER,
            parent_id INTEGER,
            component_name TEXT,
            component_type TEXT,
            label TEXT,
            part_number TEXT,
            level INTEGER DEFAULT 0,
            description TEXT
        )
    """)

    for column, definition in [
        ("project_id", "INTEGER"),
        ("parent_id", "INTEGER"),
        ("component_name", "TEXT"),
        ("component_type", "TEXT"),
        ("label", "TEXT"),
        ("part_number", "TEXT"),
        ("level", "INTEGER DEFAULT 0"),
        ("description", "TEXT")
    ]:
        add_column_if_missing(cursor, "product_structure", column, definition)

    # KEY CHARACTERISTICS
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS key_characteristics (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER,
            component_id INTEGER,
            characteristic TEXT,
            specification TEXT,
            tolerance TEXT,
            severity INTEGER,
            responsibility TEXT
        )
    """)

    for column, definition in [
        ("project_id", "INTEGER"),
        ("component_id", "INTEGER"),
        ("characteristic", "TEXT"),
        ("specification", "TEXT"),
        ("tolerance", "TEXT"),
        ("severity", "INTEGER"),
        ("responsibility", "TEXT")
    ]:
        add_column_if_missing(cursor, "key_characteristics", column, definition)

    # FUNCTIONAL LINKS
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS functional_links (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER,
            function_id INTEGER,
            component_id INTEGER,
            requirement TEXT
        )
    """)

    for column, definition in [
        ("project_id", "INTEGER"),
        ("function_id", "INTEGER"),
        ("component_id", "INTEGER"),
        ("requirement", "TEXT")
    ]:
        add_column_if_missing(cursor, "functional_links", column, definition)

    # DFMEA
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS dfmea (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER,
            component_id INTEGER,
            function TEXT,
            failure_mode TEXT,
            failure_effect TEXT,
            severity INTEGER,
            cause TEXT,
            occurrence INTEGER,
            prevention_control TEXT,
            detection_control TEXT,
            detection INTEGER,
            rpn INTEGER,
            recommended_action TEXT,
            responsibility TEXT,
            target_date TEXT,
            action_status TEXT
        )
    """)

    for column, definition in [
        ("project_id", "INTEGER"),
        ("component_id", "INTEGER"),
        ("function", "TEXT"),
        ("failure_mode", "TEXT"),
        ("failure_effect", "TEXT"),
        ("severity", "INTEGER"),
        ("cause", "TEXT"),
        ("occurrence", "INTEGER"),
        ("prevention_control", "TEXT"),
        ("detection_control", "TEXT"),
        ("detection", "INTEGER"),
        ("rpn", "INTEGER"),
        ("recommended_action", "TEXT"),
        ("responsibility", "TEXT"),
        ("target_date", "TEXT"),
        ("action_status", "TEXT")
    ]:
        add_column_if_missing(cursor, "dfmea", column, definition)

    # PFMEA
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS pfmea (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER,
            component_id INTEGER,
            process_step TEXT,
            process_function TEXT,
            failure_mode TEXT,
            failure_effect TEXT,
            severity INTEGER,
            cause TEXT,
            occurrence INTEGER,
            prevention_control TEXT,
            detection_control TEXT,
            detection INTEGER,
            rpn INTEGER,
            recommended_action TEXT,
            responsibility TEXT,
            target_date TEXT,
            action_status TEXT
        )
    """)

    for column, definition in [
        ("project_id", "INTEGER"),
        ("component_id", "INTEGER"),
        ("process_step", "TEXT"),
        ("process_function", "TEXT"),
        ("failure_mode", "TEXT"),
        ("failure_effect", "TEXT"),
        ("severity", "INTEGER"),
        ("cause", "TEXT"),
        ("occurrence", "INTEGER"),
        ("prevention_control", "TEXT"),
        ("detection_control", "TEXT"),
        ("detection", "INTEGER"),
        ("rpn", "INTEGER"),
        ("recommended_action", "TEXT"),
        ("responsibility", "TEXT"),
        ("target_date", "TEXT"),
        ("action_status", "TEXT")
    ]:
        add_column_if_missing(cursor, "pfmea", column, definition)

    # CONTROL PLAN
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS control_plan (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER,
            component_id INTEGER,
            process_step TEXT,
            characteristic TEXT,
            specification TEXT,
            control_method TEXT,
            measurement_method TEXT,
            sample_size TEXT,
            frequency TEXT,
            responsibility TEXT,
            reaction_plan TEXT
        )
    """)

    for column, definition in [
        ("project_id", "INTEGER"),
        ("component_id", "INTEGER"),
        ("process_step", "TEXT"),
        ("characteristic", "TEXT"),
        ("specification", "TEXT"),
        ("control_method", "TEXT"),
        ("measurement_method", "TEXT"),
        ("sample_size", "TEXT"),
        ("frequency", "TEXT"),
        ("responsibility", "TEXT"),
        ("reaction_plan", "TEXT")
    ]:
        add_column_if_missing(cursor, "control_plan", column, definition)

    conn.commit()
    conn.close()


# =========================================================
# LOGIN
# =========================================================

@app.before_request
def require_login():
    # Public pages
    if request.endpoint in {
        "login",
        "register",
        "forgot_password",
        "static"
    }:
        return None

    # All other pages require login
    if "user_id" not in session:
        return redirect(url_for("login"))
    
@app.route("/login", methods=["GET", "POST"])
def login():
    if session.get("user_id"):
        return redirect(url_for("dashboard"))

    error = None

    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")

        conn = get_db()
        user = conn.execute(
            "SELECT * FROM users WHERE username = ?",
            (username,)
        ).fetchone()
        conn.close()

        if user and check_password_hash(user["password_hash"], password):
            session.clear()
            session["user_id"] = user["id"]
            session["username"] = user["username"]
            session["name"] = user["name"] if "name" in user.keys() else ""
            session["organisation"] = user["organisation"] if "organisation" in user.keys() else ""
            session["language"] = user["language"] if "language" in user.keys() else "English"
            return redirect(url_for("dashboard"))

        error = "Invalid username or password."

    return render_template("login.html", error=error)


@app.route("/register", methods=["GET", "POST"])
def register():
    if session.get("user_id"):
        return redirect(url_for("dashboard"))

    error = None

    if request.method == "POST":
        name = request.form.get("name", "").strip()
        organisation = request.form.get("organisation", "").strip()
        language = request.form.get("language", "English").strip()
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        confirm_password = request.form.get("confirm_password", "")

        if not name or not organisation or not username or not password:
            error = "Please fill in all required fields."
            return render_template("register.html", error=error)

        if organisation.lower() != "antolin":
            error = "Account creation is allowed only for Antolin organisation."
            return render_template("register.html", error=error)

        if password != confirm_password:
            error = "Passwords do not match."
            return render_template("register.html", error=error)

        if len(password) < 8:
            error = "Password must contain at least 8 characters."
            return render_template("register.html", error=error)

        if not re.search(r"[A-Z]", password):
            error = "Password must contain at least 1 uppercase letter."
            return render_template("register.html", error=error)

        if not re.search(r"\d", password):
            error = "Password must contain at least 1 number."
            return render_template("register.html", error=error)

        if not re.search(r"[^A-Za-z0-9]", password):
            error = "Password must contain at least 1 special symbol."
            return render_template("register.html", error=error)

        allowed_languages = [
            "English", "French", "Spanish", "German", "Italian", "Portuguese"
        ]
        if language not in allowed_languages:
            language = "English"

        conn = get_db()
        existing_user = conn.execute(
            "SELECT id FROM users WHERE username = ?",
            (username,)
        ).fetchone()

        if existing_user:
            conn.close()
            return render_template(
                "register.html",
                error="Username already exists. Please choose another."
            )

        conn.execute("""
            INSERT INTO users
            (name, organisation, language, username, password_hash, created_date)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (
            name,
            "Antolin",
            language,
            username,
            generate_password_hash(password),
            date.today().isoformat()
        ))
        conn.commit()
        conn.close()

        flash("Account created successfully. Please log in.", "success")
        return redirect(url_for("login"))

    return render_template("register.html", error=error)


@app.route("/forgot-password", methods=["GET", "POST"])
def forgot_password():
    if session.get("user_id"):
        return redirect(url_for("dashboard"))

    error = None

    if request.method == "POST":
        username = request.form.get("username", "").strip()
        name = request.form.get("name", "").strip()
        organisation = request.form.get("organisation", "").strip()
        password = request.form.get("password", "")
        confirm_password = request.form.get("confirm_password", "")

        if not username or not name or not organisation or not password or not confirm_password:
            error = "Please fill in all fields."
            return render_template("forgot_password.html", error=error)

        if organisation.lower() != "antolin":
            error = "Password reset is available only for Antolin accounts."
            return render_template("forgot_password.html", error=error)

        if password != confirm_password:
            error = "Passwords do not match."
            return render_template("forgot_password.html", error=error)

        if len(password) < 8:
            error = "Password must contain at least 8 characters."
            return render_template("forgot_password.html", error=error)

        if not re.search(r"[A-Z]", password):
            error = "Password must contain at least 1 uppercase letter."
            return render_template("forgot_password.html", error=error)

        if not re.search(r"\d", password):
            error = "Password must contain at least 1 number."
            return render_template("forgot_password.html", error=error)

        if not re.search(r"[^A-Za-z0-9]", password):
            error = "Password must contain at least 1 special symbol."
            return render_template("forgot_password.html", error=error)

        conn = get_db()
        user = conn.execute("""
            SELECT id FROM users
            WHERE username = ? AND name = ? AND organisation = ?
        """, (username, name, "Antolin")).fetchone()

        if not user:
            conn.close()
            error = "The account details could not be verified."
            return render_template("forgot_password.html", error=error)

        conn.execute("""
            UPDATE users
            SET password_hash = ?
            WHERE id = ?
        """, (generate_password_hash(password), user["id"]))
        conn.commit()
        conn.close()

        flash("Password changed successfully. Please log in with your new password.", "success")
        return redirect(url_for("login"))

    return render_template("forgot_password.html", error=error)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


# =========================================================
# USER / PROJECT OWNERSHIP HELPERS
# =========================================================

def current_user_id():
    return session.get("user_id")


def owned_project(conn, project_id):
    if not project_id:
        return None
    return conn.execute(
        "SELECT * FROM projects WHERE id = ? AND user_id = ?",
        (project_id, current_user_id())
    ).fetchone()


def project_belongs_to_user(conn, project_id):
    return owned_project(conn, project_id) is not None


def record_belongs_to_user(conn, table_name, record_id):
    allowed_tables = {
        "functional_analysis", "boundary_diagram", "product_structure",
        "key_characteristics", "functional_links", "dfmea",
        "pfmea", "control_plan"
    }
    if table_name not in allowed_tables:
        return False
    row = conn.execute(
        f"SELECT 1 FROM {table_name} AS t JOIN projects AS p ON t.project_id = p.id "
        "WHERE t.id = ? AND p.user_id = ?",
        (record_id, current_user_id())
    ).fetchone()
    return row is not None


# =========================================================
# DASHBOARD
# =========================================================

@app.route("/")
def dashboard():
    if not session.get("user_id"):
        return redirect(url_for("register"))

    conn = get_db()
    projects = conn.execute(
        "SELECT * FROM projects WHERE user_id = ? ORDER BY id DESC",
        (current_user_id(),)
    ).fetchall()
    conn.close()

    return render_template(
        "dashboard.html",
        projects=projects
    )



@app.route("/dashboard")
def dashboard_alias():
    return redirect(url_for("dashboard"))


# =========================================================
# PROJECT
# =========================================================

@app.route("/project", methods=["GET", "POST"])
def project():
    if request.method == "POST":
        project_name = request.form.get("project_name", "").strip()
        product_name = request.form.get("product_name", "").strip()
        customer = request.form.get("customer", "").strip()
        oem_name = request.form.get("oem_name", "Generic").strip()
        compliance_mode = request.form.get(
            "compliance_mode", "AIAG-VDA 2019"
        ).strip()
        project_number = request.form.get("project_number", "").strip()
        created_date = request.form.get("created_date", "").strip()

        if not created_date:
            created_date = date.today().isoformat()

        if project_name:
            conn = get_db()
            conn.execute("""
                INSERT INTO projects
                (user_id, project_name, product_name, customer, oem_name,
                 compliance_mode, project_number, created_date)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                current_user_id(),
                project_name,
                product_name,
                customer,
                oem_name,
                compliance_mode,
                project_number,
                created_date
            ))
            conn.commit()
            conn.close()

        return redirect(url_for("project"))

    conn = get_db()
    projects = conn.execute(
        "SELECT * FROM projects WHERE user_id = ? ORDER BY id DESC",
        (current_user_id(),)
    ).fetchall()
    conn.close()

    return render_template(
        "project.html",
        projects=projects,
        today=date.today().isoformat()
    )


# =========================================================
# FUNCTIONAL ANALYSIS
# =========================================================

@app.route("/functional-analysis", methods=["GET", "POST"])
def functional_analysis():
    conn = get_db()

    if request.method == "POST":
        project_id = request.form.get("project_id", "")
        if project_id and not project_belongs_to_user(conn, project_id):
            conn.close()
            return "Project not found.", 404
        function = request.form.get("function", "").strip()
        requirement = request.form.get("requirement", "").strip()

        if project_id and function:
            conn.execute("""
                INSERT INTO functional_analysis
                (project_id, function, requirement)
                VALUES (?, ?, ?)
            """, (project_id, function, requirement))
            conn.commit()

    projects = conn.execute("""
        SELECT * FROM projects WHERE user_id = ? ORDER BY project_name
    """, (current_user_id(),)).fetchall()

    records = conn.execute("""
        SELECT fa.id, fa.project_id, p.project_name,
               fa.function, fa.requirement
        FROM functional_analysis AS fa
        LEFT JOIN projects AS p ON fa.project_id = p.id
        WHERE p.user_id = ?
        ORDER BY fa.id DESC
    """, (current_user_id(),)).fetchall()

    conn.close()

    return render_template(
        "functional_analysis.html",
        projects=projects,
        records=records
    )


# =========================================================
# BOUNDARY DIAGRAM
# =========================================================

@app.route("/boundary-diagram", methods=["GET", "POST"])
def boundary_diagram():
    conn = get_db()

    if request.method == "POST":
        project_id = request.form.get("project_id", "")
        if project_id and not project_belongs_to_user(conn, project_id):
            conn.close()
            return "Project not found.", 404
        external_element = request.form.get(
            "external_element", ""
        ).strip()
        interaction = request.form.get("interaction", "").strip()
        direction = request.form.get("direction", "").strip()
        description = request.form.get("description", "").strip()

        if project_id and external_element:
            conn.execute("""
                INSERT INTO boundary_diagram
                (project_id, external_element, interaction,
                 direction, description)
                VALUES (?, ?, ?, ?, ?)
            """, (
                project_id,
                external_element,
                interaction,
                direction,
                description
            ))
            conn.commit()

    projects = conn.execute(
        "SELECT * FROM projects WHERE user_id = ? ORDER BY project_name",
        (current_user_id(),)
    ).fetchall()

    boundaries = conn.execute("""
        SELECT bd.id, bd.project_id, bd.external_element,
               bd.interaction, bd.direction, bd.description,
               p.project_name
        FROM boundary_diagram AS bd
        LEFT JOIN projects AS p ON bd.project_id = p.id
        WHERE p.user_id = ?
        ORDER BY bd.id DESC
    """, (current_user_id(),)).fetchall()

    conn.close()

    return render_template(
        "boundary_diagram.html",
        projects=projects,
        boundaries=boundaries
    )


# =========================================================
# PRODUCT STRUCTURE
# =========================================================

@app.route("/product-structure", methods=["GET", "POST"])
def product_structure():
    conn = get_db()
    selected_project_id = request.args.get("project_id", "")
    if selected_project_id and not project_belongs_to_user(conn, selected_project_id):
        selected_project_id = ""

    if request.method == "POST":
        project_id = request.form.get("project_id", "")
        if project_id and not project_belongs_to_user(conn, project_id):
            conn.close()
            return "Project not found.", 404
        parent_id = request.form.get("parent_id", "")
        component_name = request.form.get(
            "component_name", ""
        ).strip()
        component_type = request.form.get(
            "component_type", ""
        ).strip()
        label = request.form.get("label", "").strip()
        part_number = request.form.get("part_number", "").strip()
        description = request.form.get("description", "").strip()

        if parent_id == "":
            parent_id = None

        # Calculate hierarchy level from the actual parent instead of
        # trusting a value sent by the browser.
        level = 0
        if parent_id:
            parent = conn.execute("""
                SELECT id, level
                FROM product_structure
                WHERE id = ? AND project_id = ?
            """, (parent_id, project_id)).fetchone()

            if parent:
                level = int(parent["level"] or 0) + 1
            else:
                parent_id = None

        if project_id and component_name:
            conn.execute("""
                INSERT INTO product_structure
                (project_id, parent_id, component_name, component_type,
                 label, part_number, level, description)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                project_id,
                parent_id,
                component_name,
                component_type,
                label,
                part_number,
                level,
                description
            ))
            conn.commit()
            conn.close()

            return redirect(
                url_for(
                    "product_structure",
                    project_id=project_id
                )
            )

    projects = conn.execute("""
        SELECT id, project_name, product_name
        FROM projects
        WHERE user_id = ?
        ORDER BY id DESC
    """, (current_user_id(),)).fetchall()

    if selected_project_id:
        components = conn.execute("""
            SELECT * FROM product_structure
            WHERE project_id = ?
            ORDER BY level, id
        """, (selected_project_id,)).fetchall()
    else:
        components = []

    conn.close()

    return render_template(
        "product_structure.html",
        projects=projects,
        components=components,
        selected_project_id=selected_project_id
    )


# =========================================================
# KEY CHARACTERISTICS
# =========================================================


@app.route("/key-characteristics", methods=["GET", "POST"])
def key_characteristics():
    conn = get_db()
    selected_project_id = request.args.get("project_id", "")

    if selected_project_id and not project_belongs_to_user(
        conn, selected_project_id
    ):
        selected_project_id = ""

    if request.method == "POST":
        project_id = request.form.get("project_id", "").strip()
        record_id = request.form.get("record_id", "").strip()
        component_id = request.form.get("component_id", "").strip()
        characteristic = request.form.get("characteristic", "").strip()
        specification = request.form.get("specification", "").strip()
        tolerance = request.form.get("tolerance", "").strip()
        severity = request.form.get("severity", "1").strip()
        responsibility = request.form.get("responsibility", "").strip()

        if not project_id or not project_belongs_to_user(conn, project_id):
            conn.close()
            return "Project not found.", 404

        if not component_id or not characteristic:
            conn.close()
            return "Component and SC point are required.", 400

        component = conn.execute(
            """
            SELECT id
            FROM product_structure
            WHERE id = ? AND project_id = ?
            """,
            (component_id, project_id)
        ).fetchone()

        if not component:
            conn.close()
            return "Component not found in this project.", 400

        if record_id:
            # Update an existing record instead of creating a duplicate.
            existing = conn.execute(
                """
                SELECT kc.id
                FROM key_characteristics AS kc
                JOIN projects AS p ON p.id = kc.project_id
                WHERE kc.id = ?
                  AND kc.project_id = ?
                  AND p.user_id = ?
                """,
                (record_id, project_id, current_user_id())
            ).fetchone()

            if not existing:
                conn.close()
                return "Key Characteristics record not found.", 404

            conn.execute(
                """
                UPDATE key_characteristics
                SET component_id = ?,
                    characteristic = ?,
                    specification = ?,
                    tolerance = ?,
                    severity = ?,
                    responsibility = ?
                WHERE id = ? AND project_id = ?
                """,
                (
                    component_id,
                    characteristic,
                    specification,
                    tolerance,
                    severity,
                    responsibility,
                    record_id,
                    project_id
                )
            )

            conn.commit()
            conn.close()

            return {
                "success": True,
                "record_id": int(record_id),
                "message": "Record updated"
            }

        # Create a new record.
        cursor = conn.execute(
            """
            INSERT INTO key_characteristics
            (
                project_id, component_id, characteristic,
                specification, tolerance, severity, responsibility
            )
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                project_id,
                component_id,
                characteristic,
                specification,
                tolerance,
                severity,
                responsibility
            )
        )

        new_record_id = cursor.lastrowid
        conn.commit()
        conn.close()

        return {
            "success": True,
            "record_id": new_record_id,
            "message": "Record created"
        }

    projects = conn.execute(
        """
        SELECT id, project_name, product_name
        FROM projects
        WHERE user_id = ?
        ORDER BY id DESC
        """,
        (current_user_id(),)
    ).fetchall()

    if selected_project_id:
        components = conn.execute(
            """
            SELECT id, project_id, component_name, component_type,
                   label, part_number, level
            FROM product_structure
            WHERE project_id = ?
            ORDER BY level, id
            """,
            (selected_project_id,)
        ).fetchall()

        records = conn.execute(
            """
            SELECT kc.id, kc.project_id, kc.component_id,
                   p.project_name, ps.component_name,
                   kc.characteristic, kc.specification,
                   kc.tolerance, kc.severity, kc.responsibility
            FROM key_characteristics AS kc
            LEFT JOIN projects AS p ON kc.project_id = p.id
            LEFT JOIN product_structure AS ps ON kc.component_id = ps.id
            WHERE p.user_id = ? AND kc.project_id = ?
            ORDER BY kc.id ASC
            """,
            (current_user_id(), selected_project_id)
        ).fetchall()
    else:
        components = []

        records = conn.execute(
            """
            SELECT kc.id, kc.project_id, kc.component_id,
                   p.project_name, ps.component_name,
                   kc.characteristic, kc.specification,
                   kc.tolerance, kc.severity, kc.responsibility
            FROM key_characteristics AS kc
            LEFT JOIN projects AS p ON kc.project_id = p.id
            LEFT JOIN product_structure AS ps ON kc.component_id = ps.id
            WHERE p.user_id = ?
            ORDER BY kc.id ASC
            """,
            (current_user_id(),)
        ).fetchall()

    conn.close()

    return render_template(
        "key_characteristics.html",
        projects=projects,
        components=components,
        records=records,
        selected_project_id=selected_project_id
    )
# =========================================================
# FUNCTIONAL LINKS
# =========================================================

@app.route("/functional-links", methods=["GET", "POST"])
def functional_links():
    """Visual Function -> Product Structure allocation matrix.

    A click on an intersection toggles the relationship between one
    Functional Analysis record and one Product Structure component.
    The existing functional_links table is reused, so no migration is
    required for existing projects or links.
    """
    conn = get_db()
    selected_project_id = request.args.get("project_id", "")

    if selected_project_id and not project_belongs_to_user(
        conn, selected_project_id
    ):
        selected_project_id = ""

    # ---------------------------------------------------------
    # CLICK-TO-LINK API
    # ---------------------------------------------------------
    if request.method == "POST":
        data = request.get_json(silent=True)

        if data:
            project_id = str(data.get("project_id", "")).strip()
            function_id = str(data.get("function_id", "")).strip()
            component_id = str(data.get("component_id", "")).strip()

            if not project_id or not function_id or not component_id:
                conn.close()
                return {
                    "success": False,
                    "message": "Project, function and component are required."
                }, 400

            if not project_belongs_to_user(conn, project_id):
                conn.close()
                return {
                    "success": False,
                    "message": "Project not found."
                }, 404

            function_row = conn.execute("""
                SELECT id, project_id, function, requirement
                FROM functional_analysis
                WHERE id = ? AND project_id = ?
            """, (function_id, project_id)).fetchone()

            component_row = conn.execute("""
                SELECT id, project_id, component_name
                FROM product_structure
                WHERE id = ? AND project_id = ?
            """, (component_id, project_id)).fetchone()

            if not function_row or not component_row:
                conn.close()
                return {
                    "success": False,
                    "message": "Function or component not found."
                }, 404

            existing = conn.execute("""
                SELECT id
                FROM functional_links
                WHERE project_id = ?
                  AND function_id = ?
                  AND component_id = ?
            """, (project_id, function_id, component_id)).fetchone()

            if existing:
                conn.execute(
                    "DELETE FROM functional_links WHERE id = ?",
                    (existing["id"],)
                )
                linked = False
                link_id = existing["id"]
            else:
                # The Functional Analysis requirement is the traceability
                # requirement. The user no longer has to type it manually.
                requirement = (function_row["requirement"] or "").strip()
                conn.execute("""
                    INSERT INTO functional_links
                    (project_id, function_id, component_id, requirement)
                    VALUES (?, ?, ?, ?)
                """, (
                    project_id,
                    function_id,
                    component_id,
                    requirement
                ))
                link_id = conn.execute(
                    "SELECT last_insert_rowid()"
                ).fetchone()[0]
                linked = True

            conn.commit()
            conn.close()

            return {
                "success": True,
                "linked": linked,
                "link_id": link_id,
                "message": "Link created." if linked else "Link removed."
            }

        # Keep a small compatibility path for old form submissions.
        project_id = request.form.get("project_id", "").strip()
        function_id = request.form.get("function_id", "").strip()
        component_id = request.form.get("component_id", "").strip()

        if project_id and function_id and component_id:
            if not project_belongs_to_user(conn, project_id):
                conn.close()
                return "Project not found.", 404

            function_row = conn.execute("""
                SELECT id, requirement
                FROM functional_analysis
                WHERE id = ? AND project_id = ?
            """, (function_id, project_id)).fetchone()

            component_row = conn.execute("""
                SELECT id
                FROM product_structure
                WHERE id = ? AND project_id = ?
            """, (component_id, project_id)).fetchone()

            if function_row and component_row:
                existing = conn.execute("""
                    SELECT id
                    FROM functional_links
                    WHERE project_id = ?
                      AND function_id = ?
                      AND component_id = ?
                """, (project_id, function_id, component_id)).fetchone()

                if existing:
                    conn.execute(
                        "DELETE FROM functional_links WHERE id = ?",
                        (existing["id"],)
                    )
                else:
                    conn.execute("""
                        INSERT INTO functional_links
                        (project_id, function_id, component_id, requirement)
                        VALUES (?, ?, ?, ?)
                    """, (
                        project_id,
                        function_id,
                        component_id,
                        function_row["requirement"] or ""
                    ))
                conn.commit()

        conn.close()
        return redirect(
            url_for("functional_links", project_id=project_id)
        )

    # ---------------------------------------------------------
    # PROJECTS / MATRIX DATA
    # ---------------------------------------------------------
    projects = conn.execute("""
        SELECT id, project_name, product_name
        FROM projects
        WHERE user_id = ?
        ORDER BY id DESC
    """, (current_user_id(),)).fetchall()

    functions = []
    components = []
    linked_map = {}
    records = []

    if selected_project_id:
        functions = conn.execute("""
            SELECT id, project_id, function, requirement
            FROM functional_analysis
            WHERE project_id = ?
            ORDER BY id ASC
        """, (selected_project_id,)).fetchall()

        components = conn.execute("""
            SELECT id, project_id, parent_id, component_name,
                   component_type, label, part_number, level
            FROM product_structure
            WHERE project_id = ?
            ORDER BY level ASC, id ASC
        """, (selected_project_id,)).fetchall()

        links = conn.execute("""
            SELECT id, function_id, component_id
            FROM functional_links
            WHERE project_id = ?
        """, (selected_project_id,)).fetchall()

        linked_map = {
            f"{row['function_id']}:{row['component_id']}": row['id']
            for row in links
        }

    records = conn.execute("""
        SELECT fl.id, fl.project_id,
               fl.requirement AS linked_requirement,
               p.project_name,
               fa.function,
               fa.requirement AS function_requirement,
               ps.component_name
        FROM functional_links AS fl
        LEFT JOIN projects AS p ON fl.project_id = p.id
        LEFT JOIN functional_analysis AS fa ON fl.function_id = fa.id
        LEFT JOIN product_structure AS ps ON fl.component_id = ps.id
        WHERE p.user_id = ?
        ORDER BY fl.id DESC
    """, (current_user_id(),)).fetchall()

    conn.close()

    return render_template(
        "functional_links.html",
        projects=projects,
        functions=functions,
        components=components,
        linked_map=linked_map,
        records=records,
        selected_project_id=selected_project_id
    )


# =========================================================
# DFMEA
# =========================================================

@app.route("/dfmea", methods=["GET", "POST"])
def dfmea():
    conn = get_db()
    selected_project_id = request.args.get("project_id", "")
    if selected_project_id and not project_belongs_to_user(conn, selected_project_id):
        selected_project_id = ""

    if request.method == "POST":
        project_id = request.form.get("project_id", "")
        if project_id and not project_belongs_to_user(conn, project_id):
            conn.close()
            return "Project not found.", 404
        component_id = request.form.get("component_id", "")
        function = request.form.get("function", "").strip()
        failure_mode = request.form.get(
            "failure_mode", ""
        ).strip()
        failure_effect = request.form.get(
            "failure_effect", ""
        ).strip()
        severity = request.form.get("severity", "1")
        cause = request.form.get("cause", "").strip()
        occurrence = request.form.get("occurrence", "1")
        prevention_control = request.form.get(
            "prevention_control", ""
        ).strip()
        detection_control = request.form.get(
            "detection_control", ""
        ).strip()
        detection = request.form.get("detection", "1")
        recommended_action = request.form.get(
            "recommended_action", ""
        ).strip()
        responsibility = request.form.get(
            "responsibility", ""
        ).strip()
        target_date = request.form.get(
            "target_date", ""
        ).strip()
        action_status = request.form.get(
            "action_status", "Open"
        ).strip()

        try:
            s = max(1, min(10, int(severity)))
            o = max(1, min(10, int(occurrence)))
            d = max(1, min(10, int(detection)))
            rpn = s * o * d
        except ValueError:
            s = o = d = rpn = 1

        if project_id and failure_mode:
            conn.execute("""
                INSERT INTO dfmea
                (project_id, component_id, function, failure_mode,
                 failure_effect, severity, cause, occurrence,
                 prevention_control, detection_control, detection,
                 rpn, recommended_action, responsibility,
                 target_date, action_status)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                project_id,
                component_id or None,
                function,
                failure_mode,
                failure_effect,
                s,
                cause,
                o,
                prevention_control,
                detection_control,
                d,
                rpn,
                recommended_action,
                responsibility,
                target_date,
                action_status
            ))
            conn.commit()
            conn.close()

            return redirect(
                url_for("dfmea", project_id=project_id)
            )

    projects = conn.execute("""
        SELECT id, project_name, product_name
        FROM projects
        WHERE user_id = ?
        ORDER BY id DESC
    """, (current_user_id(),)).fetchall()

    if selected_project_id:
        components = conn.execute("""
            SELECT id, component_name, component_type,
                   part_number, level
            FROM product_structure
            WHERE project_id = ?
            ORDER BY level, id
        """, (selected_project_id,)).fetchall()
    else:
        components = []

    records = conn.execute("""
        SELECT d.*, p.project_name, ps.component_name
        FROM dfmea AS d
        LEFT JOIN projects AS p ON d.project_id = p.id
        LEFT JOIN product_structure AS ps ON d.component_id = ps.id
        ORDER BY d.id DESC
    """).fetchall()

    conn.close()

    return render_template(
        "dfmea.html",
        projects=projects,
        components=components,
        records=records,
        selected_project_id=selected_project_id
    )


# =========================================================
# PFMEA
# =========================================================

@app.route("/pfmea", methods=["GET", "POST"])
def pfmea():
    conn = get_db()
    selected_project_id = request.args.get("project_id", "")
    if selected_project_id and not project_belongs_to_user(conn, selected_project_id):
        selected_project_id = ""

    if request.method == "POST":
        project_id = request.form.get("project_id", "")
        if project_id and not project_belongs_to_user(conn, project_id):
            conn.close()
            return "Project not found.", 404
        component_id = request.form.get("component_id", "")
        process_step = request.form.get(
            "process_step", ""
        ).strip()
        process_function = request.form.get(
            "process_function", ""
        ).strip()
        failure_mode = request.form.get(
            "failure_mode", ""
        ).strip()
        failure_effect = request.form.get(
            "failure_effect", ""
        ).strip()
        severity = request.form.get("severity", "1")
        cause = request.form.get("cause", "").strip()
        occurrence = request.form.get("occurrence", "1")
        prevention_control = request.form.get(
            "prevention_control", ""
        ).strip()
        detection_control = request.form.get(
            "detection_control", ""
        ).strip()
        detection = request.form.get("detection", "1")
        recommended_action = request.form.get(
            "recommended_action", ""
        ).strip()
        responsibility = request.form.get(
            "responsibility", ""
        ).strip()
        target_date = request.form.get(
            "target_date", ""
        ).strip()
        action_status = request.form.get(
            "action_status", "Open"
        ).strip()

        try:
            s = max(1, min(10, int(severity)))
            o = max(1, min(10, int(occurrence)))
            d = max(1, min(10, int(detection)))
            rpn = s * o * d
        except ValueError:
            s = o = d = rpn = 1

        if project_id and failure_mode:
            conn.execute("""
                INSERT INTO pfmea
                (project_id, component_id, process_step,
                 process_function, failure_mode, failure_effect,
                 severity, cause, occurrence, prevention_control,
                 detection_control, detection, rpn,
                 recommended_action, responsibility,
                 target_date, action_status)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                project_id,
                component_id or None,
                process_step,
                process_function,
                failure_mode,
                failure_effect,
                s,
                cause,
                o,
                prevention_control,
                detection_control,
                d,
                rpn,
                recommended_action,
                responsibility,
                target_date,
                action_status
            ))
            conn.commit()
            conn.close()

            return redirect(
                url_for("pfmea", project_id=project_id)
            )

    projects = conn.execute("""
        SELECT id, project_name, product_name
        FROM projects
        WHERE user_id = ?
        ORDER BY id DESC
    """, (current_user_id(),)).fetchall()

    if selected_project_id:
        components = conn.execute("""
            SELECT id, component_name, component_type,
                   part_number, level
            FROM product_structure
            WHERE project_id = ?
            ORDER BY level, id
        """, (selected_project_id,)).fetchall()
    else:
        components = []

    records = conn.execute("""
        SELECT p.*, pr.project_name, ps.component_name
        FROM pfmea AS p
        LEFT JOIN projects AS pr ON p.project_id = pr.id
        LEFT JOIN product_structure AS ps ON p.component_id = ps.id
        ORDER BY p.id DESC
    """).fetchall()

    conn.close()

    return render_template(
        "pfmea.html",
        projects=projects,
        components=components,
        records=records,
        selected_project_id=selected_project_id
    )

# =========================================================
# CONTROL PLAN
# =========================================================

@app.route("/control-plan", methods=["GET", "POST"])
def control_plan():
    conn = get_db()
    selected_project_id = request.args.get("project_id", "")
    if selected_project_id and not project_belongs_to_user(conn, selected_project_id):
        selected_project_id = ""

    if request.method == "POST":
        project_id = request.form.get("project_id", "")
        if project_id and not project_belongs_to_user(conn, project_id):
            conn.close()
            return "Project not found.", 404
        component_id = request.form.get("component_id", "")
        process_step = request.form.get(
            "process_step", ""
        ).strip()
        characteristic = request.form.get(
            "characteristic", ""
        ).strip()
        specification = request.form.get(
            "specification", ""
        ).strip()
        control_method = request.form.get(
            "control_method", ""
        ).strip()
        measurement_method = request.form.get(
            "measurement_method", ""
        ).strip()
        sample_size = request.form.get(
            "sample_size", ""
        ).strip()
        frequency = request.form.get(
            "frequency", ""
        ).strip()
        responsibility = request.form.get(
            "responsibility", ""
        ).strip()
        reaction_plan = request.form.get(
            "reaction_plan", ""
        ).strip()

        if project_id and characteristic:
            conn.execute("""
                INSERT INTO control_plan
                (project_id, component_id, process_step,
                 characteristic, specification, control_method,
                 measurement_method, sample_size, frequency,
                 responsibility, reaction_plan)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                project_id,
                component_id or None,
                process_step,
                characteristic,
                specification,
                control_method,
                measurement_method,
                sample_size,
                frequency,
                responsibility,
                reaction_plan
            ))
            conn.commit()
            conn.close()

            return redirect(
                url_for(
                    "control_plan",
                    project_id=project_id
                )
            )

    projects = conn.execute("""
        SELECT id, project_name, product_name
        FROM projects
        WHERE user_id = ?
        ORDER BY id DESC
    """, (current_user_id(),)).fetchall()

    if selected_project_id:
        components = conn.execute("""
            SELECT id, component_name, component_type,
                   part_number, level
            FROM product_structure
            WHERE project_id = ?
            ORDER BY level, id
        """, (selected_project_id,)).fetchall()
    else:
        components = []

    records = conn.execute("""
        SELECT cp.*, p.project_name, ps.component_name
        FROM control_plan AS cp
        LEFT JOIN projects AS p ON cp.project_id = p.id
        LEFT JOIN product_structure AS ps ON cp.component_id = ps.id
        ORDER BY cp.id DESC
    """).fetchall()

    conn.close()

    return render_template(
        "control_plan.html",
        projects=projects,
        components=components,
        records=records,
        selected_project_id=selected_project_id
    )



# =========================================================
# EDIT ROUTES
# =========================================================

@app.route("/functional-analysis/edit/<int:record_id>", methods=["GET", "POST"])
def edit_functional_analysis(record_id):
    conn = get_db()
    record = conn.execute(
        """SELECT t.* FROM functional_analysis AS t
           JOIN projects AS p ON t.project_id = p.id
           WHERE t.id = ? AND p.user_id = ?""",
        (record_id, current_user_id())
    ).fetchone()

    if record is None:
        conn.close()
        return "Functional analysis record not found.", 404

    if request.method == "POST":
        function = request.form.get("function", "").strip()
        requirement = request.form.get("requirement", "").strip()

        if function:
            conn.execute("""
                UPDATE functional_analysis
                SET function = ?, requirement = ?
                WHERE id = ?
            """, (function, requirement, record_id))
            conn.commit()

        conn.close()
        return redirect(url_for("functional_analysis"))

    conn.close()
    return render_template(
        "edit_functional_analysis.html",
        record=record
    )


@app.route("/key-characteristics/edit/<int:record_id>", methods=["GET", "POST"])
def edit_key_characteristics(record_id):
    conn = get_db()
    record = conn.execute(
        """SELECT t.* FROM key_characteristics AS t
           JOIN projects AS p ON t.project_id = p.id
           WHERE t.id = ? AND p.user_id = ?""",
        (record_id, current_user_id())
    ).fetchone()

    if record is None:
        conn.close()
        return "Key characteristic record not found.", 404

    if request.method == "POST":
        project_id = request.form.get("project_id", "")
        if project_id and not project_belongs_to_user(conn, project_id):
            conn.close()
            return "Project not found.", 404
        component_id = request.form.get("component_id", "")
        characteristic = request.form.get("characteristic", "").strip()
        specification = request.form.get("specification", "").strip()
        tolerance = request.form.get("tolerance", "").strip()
        severity = request.form.get("severity", "1")
        responsibility = request.form.get("responsibility", "").strip()

        conn.execute("""
            UPDATE key_characteristics
            SET project_id = ?, component_id = ?, characteristic = ?,
                specification = ?, tolerance = ?, severity = ?,
                responsibility = ?
            WHERE id = ?
        """, (
            project_id,
            component_id or None,
            characteristic,
            specification,
            tolerance,
            severity,
            responsibility,
            record_id
        ))
        conn.commit()
        conn.close()
        return redirect(url_for("key_characteristics"))

    projects = conn.execute("""
        SELECT id, project_name, product_name
        FROM projects
        WHERE user_id = ?
        ORDER BY id DESC
    """, (current_user_id(),)).fetchall()

    components = conn.execute("""
        SELECT id, project_id, component_name, component_type,
               label, part_number, level
        FROM product_structure
        WHERE project_id = ?
        ORDER BY level, id
    """, (record["project_id"],)).fetchall()

    conn.close()

    return render_template(
        "edit_key_characteristics.html",
        record=record,
        projects=projects,
        components=components
    )


@app.route("/functional-links/edit/<int:record_id>", methods=["GET", "POST"])
def edit_functional_links(record_id):
    conn = get_db()
    record = conn.execute(
        """SELECT t.* FROM functional_links AS t
           JOIN projects AS p ON t.project_id = p.id
           WHERE t.id = ? AND p.user_id = ?""",
        (record_id, current_user_id())
    ).fetchone()

    if record is None:
        conn.close()
        return "Functional link record not found.", 404

    if request.method == "POST":
        project_id = request.form.get("project_id", "")
        if project_id and not project_belongs_to_user(conn, project_id):
            conn.close()
            return "Project not found.", 404
        function_id = request.form.get("function_id", "")
        component_id = request.form.get("component_id", "")
        requirement = request.form.get("requirement", "").strip()

        conn.execute("""
            UPDATE functional_links
            SET project_id = ?, function_id = ?, component_id = ?,
                requirement = ?
            WHERE id = ?
        """, (
            project_id,
            function_id or None,
            component_id or None,
            requirement,
            record_id
        ))
        conn.commit()
        conn.close()
        return redirect(url_for("functional_links"))

    projects = conn.execute("""
        SELECT id, project_name, product_name
        FROM projects
        WHERE user_id = ?
        ORDER BY id DESC
    """, (current_user_id(),)).fetchall()

    functions = conn.execute("""
        SELECT *
        FROM functional_analysis
        WHERE project_id = ?
        ORDER BY id
    """, (record["project_id"],)).fetchall()

    components = conn.execute("""
        SELECT *
        FROM product_structure
        WHERE project_id = ?
        ORDER BY level, id
    """, (record["project_id"],)).fetchall()

    conn.close()

    return render_template(
        "edit_functional_links.html",
        record=record,
        projects=projects,
        functions=functions,
        components=components
    )


@app.route("/dfmea/edit/<int:record_id>", methods=["GET", "POST"])
def edit_dfmea(record_id):
    conn = get_db()
    record = conn.execute(
        """SELECT t.* FROM dfmea AS t
           JOIN projects AS p ON t.project_id = p.id
           WHERE t.id = ? AND p.user_id = ?""",
        (record_id, current_user_id())
    ).fetchone()

    if record is None:
        conn.close()
        return "DFMEA record not found.", 404

    if request.method == "POST":
        project_id = request.form.get("project_id", "")
        if project_id and not project_belongs_to_user(conn, project_id):
            conn.close()
            return "Project not found.", 404
        component_id = request.form.get("component_id", "")
        function = request.form.get("function", "").strip()
        failure_mode = request.form.get("failure_mode", "").strip()
        failure_effect = request.form.get("failure_effect", "").strip()
        cause = request.form.get("cause", "").strip()
        prevention_control = request.form.get("prevention_control", "").strip()
        detection_control = request.form.get("detection_control", "").strip()
        recommended_action = request.form.get("recommended_action", "").strip()
        responsibility = request.form.get("responsibility", "").strip()
        target_date = request.form.get("target_date", "").strip()
        action_status = request.form.get("action_status", "Open").strip()

        try:
            severity = max(1, min(10, int(request.form.get("severity", "1"))))
            occurrence = max(1, min(10, int(request.form.get("occurrence", "1"))))
            detection = max(1, min(10, int(request.form.get("detection", "1"))))
        except ValueError:
            severity = occurrence = detection = 1

        rpn = severity * occurrence * detection

        conn.execute("""
            UPDATE dfmea
            SET project_id = ?, component_id = ?, function = ?,
                failure_mode = ?, failure_effect = ?, severity = ?,
                cause = ?, occurrence = ?, prevention_control = ?,
                detection_control = ?, detection = ?, rpn = ?,
                recommended_action = ?, responsibility = ?,
                target_date = ?, action_status = ?
            WHERE id = ?
        """, (
            project_id,
            component_id or None,
            function,
            failure_mode,
            failure_effect,
            severity,
            cause,
            occurrence,
            prevention_control,
            detection_control,
            detection,
            rpn,
            recommended_action,
            responsibility,
            target_date,
            action_status,
            record_id
        ))
        conn.commit()
        conn.close()
        return redirect(url_for("dfmea"))

    projects = conn.execute("""
        SELECT id, project_name, product_name
        FROM projects
        WHERE user_id = ?
        ORDER BY id DESC
    """, (current_user_id(),)).fetchall()

    components = conn.execute("""
        SELECT id, component_name, component_type,
               part_number, level
        FROM product_structure
        WHERE project_id = ?
        ORDER BY level, id
    """, (record["project_id"],)).fetchall()

    conn.close()

    return render_template(
        "edit_dfmea.html",
        record=record,
        projects=projects,
        components=components
    )


@app.route("/pfmea/edit/<int:record_id>", methods=["GET", "POST"])
def edit_pfmea(record_id):
    conn = get_db()
    record = conn.execute(
        """SELECT t.* FROM pfmea AS t
           JOIN projects AS p ON t.project_id = p.id
           WHERE t.id = ? AND p.user_id = ?""",
        (record_id, current_user_id())
    ).fetchone()

    if record is None:
        conn.close()
        return "PFMEA record not found.", 404

    if request.method == "POST":
        project_id = request.form.get("project_id", "")
        if project_id and not project_belongs_to_user(conn, project_id):
            conn.close()
            return "Project not found.", 404
        component_id = request.form.get("component_id", "")
        process_step = request.form.get("process_step", "").strip()
        process_function = request.form.get("process_function", "").strip()
        failure_mode = request.form.get("failure_mode", "").strip()
        failure_effect = request.form.get("failure_effect", "").strip()
        cause = request.form.get("cause", "").strip()
        prevention_control = request.form.get("prevention_control", "").strip()
        detection_control = request.form.get("detection_control", "").strip()
        recommended_action = request.form.get("recommended_action", "").strip()
        responsibility = request.form.get("responsibility", "").strip()
        target_date = request.form.get("target_date", "").strip()
        action_status = request.form.get("action_status", "Open").strip()

        try:
            severity = max(1, min(10, int(request.form.get("severity", "1"))))
            occurrence = max(1, min(10, int(request.form.get("occurrence", "1"))))
            detection = max(1, min(10, int(request.form.get("detection", "1"))))
        except ValueError:
            severity = occurrence = detection = 1

        rpn = severity * occurrence * detection

        conn.execute("""
            UPDATE pfmea
            SET project_id = ?, component_id = ?, process_step = ?,
                process_function = ?, failure_mode = ?, failure_effect = ?,
                severity = ?, cause = ?, occurrence = ?,
                prevention_control = ?, detection_control = ?,
                detection = ?, rpn = ?, recommended_action = ?,
                responsibility = ?, target_date = ?, action_status = ?
            WHERE id = ?
        """, (
            project_id,
            component_id or None,
            process_step,
            process_function,
            failure_mode,
            failure_effect,
            severity,
            cause,
            occurrence,
            prevention_control,
            detection_control,
            detection,
            rpn,
            recommended_action,
            responsibility,
            target_date,
            action_status,
            record_id
        ))
        conn.commit()
        conn.close()
        return redirect(url_for("pfmea"))

    projects = conn.execute("""
        SELECT id, project_name, product_name
        FROM projects
        WHERE user_id = ?
        ORDER BY id DESC
    """, (current_user_id(),)).fetchall()

    components = conn.execute("""
        SELECT id, component_name, component_type,
               part_number, level
        FROM product_structure
        WHERE project_id = ?
        ORDER BY level, id
    """, (record["project_id"],)).fetchall()

    conn.close()

    return render_template(
        "edit_pfmea.html",
        record=record,
        projects=projects,
        components=components
    )


@app.route("/control-plan/edit/<int:record_id>", methods=["GET", "POST"])
def edit_control_plan(record_id):
    conn = get_db()
    record = conn.execute(
        """SELECT t.* FROM control_plan AS t
           JOIN projects AS p ON t.project_id = p.id
           WHERE t.id = ? AND p.user_id = ?""",
        (record_id, current_user_id())
    ).fetchone()

    if record is None:
        conn.close()
        return "Control Plan record not found.", 404

    if request.method == "POST":
        project_id = request.form.get("project_id", "")
        if project_id and not project_belongs_to_user(conn, project_id):
            conn.close()
            return "Project not found.", 404
        component_id = request.form.get("component_id", "")
        process_step = request.form.get("process_step", "").strip()
        characteristic = request.form.get("characteristic", "").strip()
        specification = request.form.get("specification", "").strip()
        control_method = request.form.get("control_method", "").strip()
        measurement_method = request.form.get("measurement_method", "").strip()
        sample_size = request.form.get("sample_size", "").strip()
        frequency = request.form.get("frequency", "").strip()
        responsibility = request.form.get("responsibility", "").strip()
        reaction_plan = request.form.get("reaction_plan", "").strip()

        conn.execute("""
            UPDATE control_plan
            SET project_id = ?, component_id = ?, process_step = ?,
                characteristic = ?, specification = ?, control_method = ?,
                measurement_method = ?, sample_size = ?, frequency = ?,
                responsibility = ?, reaction_plan = ?
            WHERE id = ?
        """, (
            project_id,
            component_id or None,
            process_step,
            characteristic,
            specification,
            control_method,
            measurement_method,
            sample_size,
            frequency,
            responsibility,
            reaction_plan,
            record_id
        ))
        conn.commit()
        conn.close()
        return redirect(url_for("control_plan"))

    projects = conn.execute("""
        SELECT id, project_name, product_name
        FROM projects
        WHERE user_id = ?
        ORDER BY id DESC
    """, (current_user_id(),)).fetchall()

    components = conn.execute("""
        SELECT id, component_name, component_type,
               part_number, level
        FROM product_structure
        WHERE project_id = ?
        ORDER BY level, id
    """, (record["project_id"],)).fetchall()

    conn.close()

    return render_template(
        "edit_control_plan.html",
        record=record,
        projects=projects,
        components=components
    )


@app.route("/product-structure/edit/<int:record_id>", methods=["GET", "POST"])
def edit_product_structure(record_id):
    conn = get_db()
    record = conn.execute(
        """SELECT t.* FROM product_structure AS t
           JOIN projects AS p ON t.project_id = p.id
           WHERE t.id = ? AND p.user_id = ?""",
        (record_id, current_user_id())
    ).fetchone()

    if record is None:
        conn.close()
        return "Product structure record not found.", 404

    if request.method == "POST":
        project_id = request.form.get("project_id", "")
        if project_id and not project_belongs_to_user(conn, project_id):
            conn.close()
            return "Project not found.", 404
        parent_id = request.form.get("parent_id", "")
        component_name = request.form.get("component_name", "").strip()
        component_type = request.form.get("component_type", "").strip()
        label = request.form.get("label", "").strip()
        part_number = request.form.get("part_number", "").strip()
        description = request.form.get("description", "").strip()

        if parent_id == "" or parent_id == str(record_id):
            parent_id = None

        parent_level = -1
        if parent_id:
            parent = conn.execute("""
                SELECT level
                FROM product_structure
                WHERE id = ? AND project_id = ?
            """, (parent_id, project_id)).fetchone()
            if parent:
                parent_level = parent["level"]

        level = parent_level + 1

        conn.execute("""
            UPDATE product_structure
            SET project_id = ?, parent_id = ?, component_name = ?,
                component_type = ?, label = ?, part_number = ?,
                level = ?, description = ?
            WHERE id = ?
        """, (
            project_id,
            parent_id,
            component_name,
            component_type,
            label,
            part_number,
            level,
            description,
            record_id
        ))
        conn.commit()
        conn.close()

        return redirect(
            url_for("product_structure", project_id=project_id)
        )

    projects = conn.execute("""
        SELECT id, project_name, product_name
        FROM projects
        WHERE user_id = ?
        ORDER BY id DESC
    """, (current_user_id(),)).fetchall()

    components = conn.execute("""
        SELECT *
        FROM product_structure
        WHERE project_id = ?
        ORDER BY level, id
    """, (record["project_id"],)).fetchall()

    conn.close()

    return render_template(
        "edit_product_structure.html",
        record=record,
        projects=projects,
        components=components
    )



# =========================================================
# REPORTS
# =========================================================

@app.route("/reports", methods=["GET"])
def reports():
    conn = get_db()
    selected_project_id = request.args.get("project_id", "")
    if selected_project_id and not project_belongs_to_user(conn, selected_project_id):
        selected_project_id = ""

    projects = conn.execute(
        "SELECT * FROM projects WHERE user_id = ? ORDER BY id DESC",
        (current_user_id(),)
    ).fetchall()

    project_info = None

    summary = {
        "functional_analysis": 0,
        "boundary_diagram": 0,
        "product_structure": 0,
        "key_characteristics": 0,
        "functional_links": 0,
        "dfmea": 0,
        "pfmea": 0,
        "control_plan": 0
    }

    components = []
    dfmea_records = []
    pfmea_records = []

    if selected_project_id:
        project_info = conn.execute("""
            SELECT * FROM projects WHERE id = ? AND user_id = ?
        """, (selected_project_id, current_user_id())).fetchone()

        tables = [
            "functional_analysis",
            "boundary_diagram",
            "product_structure",
            "key_characteristics",
            "functional_links",
            "dfmea",
            "pfmea",
            "control_plan"
        ]

        for table in tables:
            summary[table] = conn.execute(
                f"SELECT COUNT(*) FROM {table} WHERE project_id = ?",
                (selected_project_id,)
            ).fetchone()[0]

        components = conn.execute("""
            SELECT id, component_name, component_type,
                   label, part_number, level
            FROM product_structure
            WHERE project_id = ?
            ORDER BY level, id
        """, (selected_project_id,)).fetchall()

        dfmea_records = conn.execute("""
            SELECT d.*, ps.component_name
            FROM dfmea AS d
            LEFT JOIN product_structure AS ps
              ON d.component_id = ps.id
            WHERE d.project_id = ?
            ORDER BY d.id DESC
        """, (selected_project_id,)).fetchall()

        pfmea_records = conn.execute("""
            SELECT p.*, ps.component_name
            FROM pfmea AS p
            LEFT JOIN product_structure AS ps
              ON p.component_id = ps.id
            WHERE p.project_id = ?
            ORDER BY p.id DESC
        """, (selected_project_id,)).fetchall()

    conn.close()

    return render_template(
        "reports.html",
        projects=projects,
        project_info=project_info,
        summary=summary,
        components=components,
        dfmea_records=dfmea_records,
        pfmea_records=pfmea_records,
        selected_project_id=selected_project_id
    )


# =========================================================
# PROFESSIONAL EXCEL HELPERS
# =========================================================

def apply_sheet_format(
    sheet,
    title,
    subtitle,
    landscape=True,
    tab_color="2F75B5"
):
    dark_blue = "17365D"
    medium_blue = "2F75B5"
    border_color = "B7C9D6"

    thin = Side(
        style="thin",
        color=border_color
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin
    )

    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = tab_color

    # Title row
    max_col = max(sheet.max_column, 1)

    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=max_col
    )

    title_cell = sheet.cell(1, 1)
    title_cell.value = title
    title_cell.font = Font(
        name="Aptos",
        size=18,
        bold=True,
        color="FFFFFF"
    )
    title_cell.fill = PatternFill(
        "solid",
        fgColor=dark_blue
    )
    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )
    sheet.row_dimensions[1].height = 32

    # Subtitle
    sheet.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=max_col
    )

    subtitle_cell = sheet.cell(2, 1)
    subtitle_cell.value = subtitle
    subtitle_cell.font = Font(
        name="Aptos",
        size=10,
        italic=True
    )
    subtitle_cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )
    sheet.row_dimensions[2].height = 22

    # Blank spacer row 3
    sheet.row_dimensions[3].height = 8

    # Table header row 4
    for cell in sheet[4]:
        cell.font = Font(
            name="Aptos",
            size=10,
            bold=True,
            color="FFFFFF"
        )
        cell.fill = PatternFill(
            "solid",
            fgColor=medium_blue
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )
        cell.border = border

    sheet.row_dimensions[4].height = 34

    # Data rows
    for row_number in range(5, sheet.max_row + 1):
        sheet.row_dimensions[row_number].height = 40

        for cell in sheet[row_number]:
            cell.font = Font(
                name="Aptos",
                size=10
            )
            cell.border = border
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

            if row_number % 2 == 1:
                cell.fill = PatternFill(
                    "solid",
                    fgColor="F4F8FB"
                )

    if sheet.max_row >= 4:
        sheet.auto_filter.ref = (
            f"A4:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        )

    sheet.freeze_panes = "A5"

    # Column widths
    for column in range(1, sheet.max_column + 1):
        letter = get_column_letter(column)
        maximum = 0

        for row in range(1, sheet.max_row + 1):
            value = sheet.cell(row, column).value

            if value is not None:
                maximum = max(
                    maximum,
                    len(str(value))
                )

        sheet.column_dimensions[letter].width = min(
            max(maximum + 3, 12),
            35
        )

    # Print setup
    sheet.page_setup.orientation = (
        "landscape" if landscape else "portrait"
    )
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    sheet.page_margins = PageMargins(
        left=0.25,
        right=0.25,
        top=0.5,
        bottom=0.5,
        header=0.2,
        footer=0.2
    )

    sheet.print_title_rows = "1:4"

    sheet.oddFooter.center.text = (
        "Automotive FMEA Management System"
    )
    sheet.oddFooter.right.text = (
        "Page &P of &N"
    )


def add_rpn_rules(sheet, column_number):
    if sheet.max_row < 5:
        return

    letter = get_column_letter(column_number)
    cell_range = f"{letter}5:{letter}{sheet.max_row}"

    sheet.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="greaterThanOrEqual",
            formula=["200"],
            fill=PatternFill(
                "solid",
                fgColor="F4CCCC"
            )
        )
    )

    sheet.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="between",
            formula=["100", "199"],
            fill=PatternFill(
                "solid",
                fgColor="FFF2CC"
            )
        )
    )

    sheet.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="lessThan",
            formula=["100"],
            fill=PatternFill(
                "solid",
                fgColor="D9EAD3"
            )
        )
    )


# =========================================================
# EXCEL EXPORT
# =========================================================

@app.route("/export-excel")
def export_excel():
    project_id = request.args.get("project_id")

    if not project_id:
        return "Please select a project first."

    conn = get_db()

    project = conn.execute("""
        SELECT * FROM projects WHERE id = ? AND user_id = ?
    """, (project_id, current_user_id())).fetchone()

    if not project:
        conn.close()
        return "Project not found."

    workbook = Workbook()

    # -----------------------------------------------------
    # PROJECT
    # -----------------------------------------------------

    sheet = workbook.active
    sheet.title = "Project"

    sheet.cell(4, 1, "Field")
    sheet.cell(4, 2, "Project Details")

    project_rows = [
        ("Project Name", project["project_name"]),
        ("Product Name", project["product_name"]),
        ("Customer", project["customer"]),
        ("OEM / Customer Standard", project["oem_name"]),
        ("Compliance Mode", project["compliance_mode"]),
        ("Project Number", project["project_number"]),
        ("Created Date", project["created_date"])
    ]

    for row_number, (label, value) in enumerate(
        project_rows,
        start=5
    ):
        sheet.cell(row_number, 1, label)
        sheet.cell(row_number, 2, value)

    apply_sheet_format(
        sheet,
        "AUTOMOTIVE FMEA MANAGEMENT SYSTEM",
        "Project Identification & Configuration",
        landscape=False,
        tab_color="17365D"
    )

    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 50

    for row_number in range(5, 12):
        sheet.cell(row_number, 1).font = Font(
            name="Aptos",
            size=10,
            bold=True
        )
        sheet.cell(row_number, 1).fill = PatternFill(
            "solid",
            fgColor="D9EAF7"
        )

    # -----------------------------------------------------
    # PRODUCT STRUCTURE
    # -----------------------------------------------------

    sheet = workbook.create_sheet("Product Structure")

    sheet.append([
        "Level",
        "Component",
        "Type",
        "Label",
        "Part Number",
        "Description"
    ])

    rows = conn.execute("""
        SELECT level, component_name, component_type,
               label, part_number, description
        FROM product_structure
        WHERE project_id = ?
        ORDER BY level, id
    """, (project_id,)).fetchall()

    for row in rows:
        sheet.append([
            row["level"],
            row["component_name"],
            row["component_type"],
            row["label"],
            row["part_number"],
            row["description"]
        ])

    apply_sheet_format(
        sheet,
        "PRODUCT STRUCTURE",
        "Product hierarchy and component definition",
        True,
        "5B9BD5"
    )

    # -----------------------------------------------------
    # FUNCTIONAL ANALYSIS
    # -----------------------------------------------------

    sheet = workbook.create_sheet("Functional Analysis")

    sheet.append([
        "Function",
        "Requirement"
    ])

    rows = conn.execute("""
        SELECT function, requirement
        FROM functional_analysis
        WHERE project_id = ?
        ORDER BY id
    """, (project_id,)).fetchall()

    for row in rows:
        sheet.append([
            row["function"],
            row["requirement"]
        ])

    apply_sheet_format(
        sheet,
        "FUNCTIONAL ANALYSIS",
        "Functions and associated requirements",
        True,
        "70AD47"
    )

    # -----------------------------------------------------
    # BOUNDARY DIAGRAM
    # -----------------------------------------------------

    sheet = workbook.create_sheet("Boundary Diagram")

    sheet.append([
        "External Element",
        "Interaction",
        "Direction",
        "Description"
    ])

    rows = conn.execute("""
        SELECT external_element, interaction,
               direction, description
        FROM boundary_diagram
        WHERE project_id = ?
        ORDER BY id
    """, (project_id,)).fetchall()

    for row in rows:
        sheet.append([
            row["external_element"],
            row["interaction"],
            row["direction"],
            row["description"]
        ])

    apply_sheet_format(
        sheet,
        "BOUNDARY DIAGRAM",
        "System boundaries and external interactions",
        True,
        "ED7D31"
    )

    # -----------------------------------------------------
    # KEY CHARACTERISTICS
    # -----------------------------------------------------

    sheet = workbook.create_sheet("Key Characteristics")

    sheet.append([
        "Component",
        "Characteristic",
        "Specification",
        "Tolerance",
        "Severity",
        "Responsibility"
    ])

    rows = conn.execute("""
        SELECT ps.component_name,
               kc.characteristic,
               kc.specification,
               kc.tolerance,
               kc.severity,
               kc.responsibility
        FROM key_characteristics AS kc
        LEFT JOIN product_structure AS ps
          ON kc.component_id = ps.id
        WHERE kc.project_id = ?
        ORDER BY kc.id
    """, (project_id,)).fetchall()

    for row in rows:
        sheet.append([
            row["component_name"],
            row["characteristic"],
            row["specification"],
            row["tolerance"],
            row["severity"],
            row["responsibility"]
        ])

    apply_sheet_format(
        sheet,
        "KEY CHARACTERISTICS",
        "Special characteristics and specifications",
        True,
        "FFC000"
    )

    # -----------------------------------------------------
    # FUNCTIONAL LINKS
    # -----------------------------------------------------

    sheet = workbook.create_sheet("Functional Links")

    sheet.append([
        "Function",
        "Function Requirement",
        "Component",
        "Linked Requirement"
    ])

    rows = conn.execute("""
        SELECT fa.function AS function_name,
               fa.requirement AS function_requirement,
               ps.component_name,
               fl.requirement AS linked_requirement
        FROM functional_links AS fl
        LEFT JOIN functional_analysis AS fa
          ON fl.function_id = fa.id
        LEFT JOIN product_structure AS ps
          ON fl.component_id = ps.id
        WHERE fl.project_id = ?
        ORDER BY fl.id
    """, (project_id,)).fetchall()

    for row in rows:
        sheet.append([
            row["function_name"],
            row["function_requirement"],
            row["component_name"],
            row["linked_requirement"]
        ])

    apply_sheet_format(
        sheet,
        "FUNCTIONAL LINKING",
        "Traceability between functions, requirements and components",
        True,
        "A5A5A5"
    )

    # -----------------------------------------------------
    # DFMEA
    # -----------------------------------------------------

    sheet = workbook.create_sheet("DFMEA")

    sheet.append([
        "Component",
        "Function",
        "Failure Mode",
        "Failure Effect",
        "Severity",
        "Potential Cause",
        "Occurrence",
        "Prevention Control",
        "Detection Control",
        "Detection",
        "RPN",
        "Recommended Action",
        "Responsibility",
        "Target Date",
        "Action Status"
    ])

    rows = conn.execute("""
        SELECT ps.component_name,
               d.function,
               d.failure_mode,
               d.failure_effect,
               d.severity,
               d.cause,
               d.occurrence,
               d.prevention_control,
               d.detection_control,
               d.detection,
               d.rpn,
               d.recommended_action,
               d.responsibility,
               d.target_date,
               d.action_status
        FROM dfmea AS d
        LEFT JOIN product_structure AS ps
          ON d.component_id = ps.id
        WHERE d.project_id = ?
        ORDER BY d.id
    """, (project_id,)).fetchall()

    for row in rows:
        sheet.append([
            row["component_name"],
            row["function"],
            row["failure_mode"],
            row["failure_effect"],
            row["severity"],
            row["cause"],
            row["occurrence"],
            row["prevention_control"],
            row["detection_control"],
            row["detection"],
            row["rpn"],
            row["recommended_action"],
            row["responsibility"],
            row["target_date"],
            row["action_status"]
        ])

    apply_sheet_format(
        sheet,
        "DESIGN FMEA (DFMEA)",
        "Design failure mode and effects analysis",
        True,
        "4472C4"
    )

    add_rpn_rules(sheet, 11)

    # -----------------------------------------------------
    # PFMEA
    # -----------------------------------------------------

    sheet = workbook.create_sheet("PFMEA")

    sheet.append([
        "Component",
        "Process Step",
        "Process Function",
        "Failure Mode",
        "Failure Effect",
        "Severity",
        "Potential Cause",
        "Occurrence",
        "Prevention Control",
        "Detection Control",
        "Detection",
        "RPN",
        "Recommended Action",
        "Responsibility",
        "Target Date",
        "Action Status"
    ])

    rows = conn.execute("""
        SELECT ps.component_name,
               p.process_step,
               p.process_function,
               p.failure_mode,
               p.failure_effect,
               p.severity,
               p.cause,
               p.occurrence,
               p.prevention_control,
               p.detection_control,
               p.detection,
               p.rpn,
               p.recommended_action,
               p.responsibility,
               p.target_date,
               p.action_status
        FROM pfmea AS p
        LEFT JOIN product_structure AS ps
          ON p.component_id = ps.id
        WHERE p.project_id = ?
        ORDER BY p.id
    """, (project_id,)).fetchall()

    for row in rows:
        sheet.append([
            row["component_name"],
            row["process_step"],
            row["process_function"],
            row["failure_mode"],
            row["failure_effect"],
            row["severity"],
            row["cause"],
            row["occurrence"],
            row["prevention_control"],
            row["detection_control"],
            row["detection"],
            row["rpn"],
            row["recommended_action"],
            row["responsibility"],
            row["target_date"],
            row["action_status"]
        ])

    apply_sheet_format(
        sheet,
        "PROCESS FMEA (PFMEA)",
        "Process failure mode and effects analysis",
        True,
        "70AD47"
    )

    add_rpn_rules(sheet, 12)

    # -----------------------------------------------------
    # CONTROL PLAN
    # -----------------------------------------------------

    sheet = workbook.create_sheet("Control Plan")

    sheet.append([
        "Component",
        "Process Step",
        "Characteristic",
        "Specification",
        "Control Method",
        "Measurement Method",
        "Sample Size",
        "Frequency",
        "Responsibility",
        "Reaction Plan"
    ])

    rows = conn.execute("""
        SELECT ps.component_name,
               cp.process_step,
               cp.characteristic,
               cp.specification,
               cp.control_method,
               cp.measurement_method,
               cp.sample_size,
               cp.frequency,
               cp.responsibility,
               cp.reaction_plan
        FROM control_plan AS cp
        LEFT JOIN product_structure AS ps
          ON cp.component_id = ps.id
        WHERE cp.project_id = ?
        ORDER BY cp.id
    """, (project_id,)).fetchall()

    for row in rows:
        sheet.append([
            row["component_name"],
            row["process_step"],
            row["characteristic"],
            row["specification"],
            row["control_method"],
            row["measurement_method"],
            row["sample_size"],
            row["frequency"],
            row["responsibility"],
            row["reaction_plan"]
        ])

    apply_sheet_format(
        sheet,
        "CONTROL PLAN",
        "Process controls, measurements and reaction plans",
        True,
        "5B9BD5"
    )

    conn.close()

    # Save workbook
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="Automotive_FMEA_Report.xlsx",
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )


# =========================================================
# INITIALIZE
# =========================================================

setup_database()


if __name__ == "__main__":
    print("")
    print("==============================================")
    print("      AUTOMOTIVE FMEA MANAGEMENT SYSTEM")
    print("==============================================")

    app.run(
        debug=False,
        host="0.0.0.0",
        port=int(os.environ.get("PORT", 5000))
    )
